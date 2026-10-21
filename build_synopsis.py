from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = r"D:\ramanujans-magic-square\Ramanujans_Magic_Square_Project_Synopsis.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "5B6573"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_font(run, size=11, color=BLACK, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_text(p, text, **kwargs):
    r = p.add_run(text)
    set_font(r, **kwargs)
    return r


def set_keep_with_next(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    ppr.append(keep)


def add_page_number(paragraph):
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    add_text(p, text)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    add_text(p, text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    add_text(p, text, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, label, size=10, color=NAVY, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_text(p, str(value), size=9.5)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# Styles: standard_business_brief preset.
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

# Header/footer.
header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(header_p, "Ramanujan's Magic Square | Project Synopsis", size=9, color=MUTED)
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_page_number(footer_p)

# Memo-masthead opening.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(20)
p.paragraph_format.space_after = Pt(4)
add_text(p, "PROJECT SYNOPSIS", size=23, color=NAVY, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
add_text(p, "Ramanujan's Magic Square", size=15, color=MUTED, italic=True)
for label, value in [
    ("Project type", "Full-stack web application for mathematics learning and puzzle gameplay"),
    ("Technology stack", "Node.js, Express, SQLite, vanilla HTML/CSS/JavaScript"),
    ("Core domain", "Magic-square generation, interactive practice, performance tracking and learning"),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_text(p, f"{label}: ", size=11, color=BLACK, bold=True)
    add_text(p, value, size=11)
rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(12)
rule.paragraph_format.space_after = Pt(8)
ppr = rule._p.get_or_add_pPr()
borders = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "8")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), BLUE)
borders.append(bottom)
ppr.append(borders)

add_heading(doc, "1. Introduction")
add_body(doc, "Ramanujan's Magic Square is a full-stack educational web application that turns the mathematical concept of magic squares into an interactive learning and puzzle-solving experience. A magic square is an n x n grid containing the integers 1 through n² exactly once, arranged so that every row, column and both main diagonals have the same sum, known as the magic constant.")
add_body(doc, "The application is named after Srinivasa Ramanujan, who was known for constructing date-based magic squares. It combines mathematical visualization with game features: users can generate squares, observe their construction, solve scrambled boards, compare performance, complete daily challenges and learn the underlying construction methods.")

add_heading(doc, "2. Problem Statement")
add_body(doc, "Magic squares are often presented as static examples in books or classrooms. This makes it difficult for learners to see how a valid square is constructed, test their understanding, or practise the logic in an engaging way. Existing puzzle experiences also commonly lack trustworthy scoring, persistent progress tracking and clear explanations of the mathematics.")
add_body(doc, "The project addresses this gap by providing one system that can generate valid magic squares for supported orders, validate puzzle solutions on the server, provide guided learning material, and record secure user performance data for dashboards and leaderboards.")

add_heading(doc, "3. Objectives")
for item in [
    "Generate and validate magic squares for orders 3 through 10, while clearly explaining why a normal 2 x 2 magic square is impossible.",
    "Teach the magic constant formula and the construction methods used for odd, doubly-even and singly-even orders.",
    "Provide an interactive game in which a generated square is scrambled and the user restores it by swapping tiles.",
    "Maintain fair scoring through server-side session timing and solution validation rather than trusting client-submitted scores.",
    "Support user accounts, personal history, best times, achievements, daily challenges and global leaderboards.",
    "Provide privacy-conscious account controls and administrative moderation functions."
]:
    add_bullet(doc, item)

add_heading(doc, "4. Scope of the Project")
add_body(doc, "The project is intended as a web-based learning and recreational mathematics platform. Its scope includes account-based puzzle play, educational explanations, competition features and lightweight administration.")
add_heading(doc, "Included", level=2)
for item in [
    "Magic-square generation for orders 3 to 10 using appropriate algorithms; order 2 is handled as an explained mathematical exception.",
    "Animated build, scramble-and-solve gameplay, hints with time penalties, undo support and export of a completed square as an image.",
    "Daily deterministic challenges, streak tracking, per-order and daily leaderboards, personal dashboard and achievement badges.",
    "Account registration, login, password reset, optional email verification, data export, self-service account deletion and role-based administration.",
    "A Formula & Methods learning page, including a Ramanujan-style date-based 4 x 4 square generator."
]:
    add_bullet(doc, item)
add_heading(doc, "Out of Scope", level=2)
for item in [
    "Horizontal multi-server scaling and distributed session storage; the current design uses in-memory sessions suitable for a single-process deployment.",
    "Third-party OAuth/social login, payment processing, native mobile applications and advanced real-time multiplayer.",
    "A separate enterprise database server; the current implementation uses a local SQLite database."
]:
    add_bullet(doc, item)

add_heading(doc, "5. Proposed System")
add_body(doc, "The proposed system is a client-server web application. The browser presents the educational pages and interactive puzzle board. A Node.js/Express server exposes REST endpoints for authentication, game creation, solution submission, rankings, daily challenges, achievement status, administration and Ramanujan date-square generation. SQLite stores users, attempts, active game sessions and earned achievements.")
add_body(doc, "When a game begins, the server creates a random game session and stores the exact solution and start time. The client receives the playable state and sends the final solved grid. The server validates the submitted grid against its stored solution and calculates elapsed time itself. This keeps rankings meaningful even if a user modifies browser-side code.")

add_heading(doc, "6. Modules")
add_table(doc, ["Module", "Key responsibilities"], [
    ("Authentication & Profile", "Registration, login/logout, session lookup, password change/reset, optional email verification, data export and self-service deletion."),
    ("Magic Square Engine", "Generates valid squares; uses Siamese method for odd orders, complement-swap for doubly-even orders and LUX/Strachey logic for singly-even orders."),
    ("Gameplay", "Starts server-backed sessions, renders and scrambles the grid, validates progress, supports hints/undo, records moves and submits final solutions."),
    ("Daily Challenge", "Creates the same deterministic daily puzzle for all players by date and order, prevents duplicate daily attempts and maintains streaks."),
    ("Attempts & Dashboard", "Persists completed solves, lists personal history, calculates best times and displays progress."),
    ("Leaderboards", "Shows the fastest recorded solves by square order and daily challenge leaderboard results."),
    ("Achievements", "Evaluates and awards badges such as First Square, Full Spectrum, Speed Demon and Daily Devotee."),
    ("Formula & Ramanujan Tool", "Explains the magic constant and construction methods; generates a 4 x 4 date-based Ramanujan-style square."),
    ("Administration", "Lists users and attempts, moderates inappropriate attempt records and grants or revokes administrator access.")
], [2250, 7110])

add_heading(doc, "7. Fundamental Requirements")
for item in [
    "The system shall generate a correct magic square for each supported order and calculate the magic constant using M(n) = n(n² + 1) / 2.",
    "The system shall ensure that a valid solution uses each required value once and that rows, columns and both diagonals equal the expected magic constant.",
    "The system shall create authenticated game sessions and store the solution, start time, move count and hint count on the server.",
    "The system shall reject a completion submission that does not exactly match the server-recorded solution.",
    "The system shall record verified attempts, calculate personal bests and publish ranked leaderboard results.",
    "The system shall provide daily challenges, achievement evaluation and role-based administration.",
    "The system shall protect passwords through bcrypt hashing and use expiring single-use tokens for password resets and email verification."
]:
    add_bullet(doc, item)

add_heading(doc, "8. Non-Functional Requirements")
add_table(doc, ["Quality attribute", "Requirement"], [
    ("Usability", "The interface should be responsive, understandable to learners and usable on common desktop and mobile browser sizes."),
    ("Performance", "Puzzle generation, grid validation and local SQLite queries should complete quickly enough for interactive play."),
    ("Security", "Passwords must be hashed; sessions use signed HTTP-only cookies; client scores are not trusted; reset and verification tokens expire."),
    ("Reliability", "Database foreign keys and WAL journal mode support data integrity and resilient local persistence."),
    ("Maintainability", "Backend code is organised into routes, middleware, database setup, mathematics engine and feature modules; frontend assets are separated by page and shared utilities."),
    ("Compatibility", "The system runs in modern web browsers and requires Node.js 18 or later on the server."),
    ("Privacy", "Users can export their account data and delete their own accounts; administrators do not delete other users' accounts.")
], [2250, 7110])

add_heading(doc, "9. Hardware Requirements")
add_table(doc, ["Component", "Minimum requirement"], [
    ("Client device", "Desktop, laptop, tablet or smartphone capable of running a current web browser."),
    ("Client memory", "At least 2 GB RAM recommended for normal browser operation."),
    ("Server processor", "Any modern single-core or higher processor suitable for Node.js and a small SQLite-backed web application."),
    ("Server memory", "Minimum 2 GB RAM; additional memory may be needed as concurrent usage grows."),
    ("Storage", "At least 500 MB available for application files, Node.js dependencies, SQLite database and backups."),
    ("Network", "HTTP/HTTPS network connectivity between the client browser and the deployed server.")
], [2250, 7110])

add_heading(doc, "10. Software Requirements")
add_table(doc, ["Software", "Purpose / version"], [
    ("Operating system", "Windows, Linux or macOS for development or deployment."),
    ("Node.js", "Version 18 or later; runs the application server."),
    ("Express", "HTTP server framework and routing layer."),
    ("SQLite + better-sqlite3", "Embedded relational database and Node.js database driver."),
    ("bcryptjs", "Password hashing."),
    ("express-session", "Server-side login session support."),
    ("Nodemailer", "Optional SMTP delivery of verification and password-reset messages."),
    ("Frontend", "HTML5, CSS3 and vanilla JavaScript; no front-end build step is required."),
    ("Browser", "A current version of Chrome, Edge, Firefox or Safari.")
], [2250, 7110])

add_heading(doc, "11. Database Tables")
add_body(doc, "The application uses SQLite. The schema contains four main tables; the users table is extended through lightweight migrations for daily-streak and account-management fields.")
add_table(doc, ["Table", "Purpose and principal fields"], [
    ("users", "Stores account identity and access data: id (PK), username (unique), password_hash, is_admin, created_at, daily_streak, last_daily_date, email (unique when present), email_verified, verify_token, verify_token_expires, reset_token and reset_token_expires."),
    ("attempts", "Stores verified completed puzzles: id (PK), user_id (FK), square_order, moves_made, time_taken, hints_used, is_daily, challenge_date and date_attempted. Indexed for user, order and daily leaderboard queries."),
    ("game_sessions", "Stores active/finished server-backed games: id (UUID PK), user_id (FK), square_order, solved_grid, is_daily, challenge_date, started_at, moves_made, hints_used and completed_at."),
    ("user_achievements", "Stores badges earned by each user: id (PK), user_id (FK), achievement_key and earned_at. A unique constraint prevents the same badge being awarded twice to the same user.")
], [2250, 7110])
add_body(doc, "Foreign keys link attempts, game sessions and achievements to users with cascading deletion. The database enables foreign-key enforcement and write-ahead logging (WAL) to improve integrity and local reliability.")

add_heading(doc, "12. Expected Outcome")
add_body(doc, "The completed system will provide a clear, enjoyable and secure way to learn and practise magic squares. Learners will be able to understand the target sum, see how different algorithms construct valid squares, and apply that knowledge through interactive gameplay. Players will receive fair, server-validated timing, visible personal progress, daily motivation and healthy competition through leaderboards.")
add_body(doc, "From a software-engineering perspective, the outcome is a demonstrable full-stack project: a responsive frontend, REST-based backend, persistent relational database, authentication, secure validation, role-based administration and maintainable modular code. It shows how a mathematical concept can be converted into a usable educational product rather than remaining a static theory example.")

doc.core_properties.title = "Ramanujan's Magic Square - Project Synopsis"
doc.core_properties.subject = "Project synopsis"
doc.core_properties.author = "Project Team"
doc.save(OUT)
print(OUT)
